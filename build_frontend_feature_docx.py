from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_system_feature_docx import (
    DARK_BLUE,
    INK,
    MUTED,
    WARN_FILL,
    add_h,
    add_metadata_rows,
    add_note,
    add_table,
    set_font,
    style_doc,
)


OUTPUT = Path("前端已实现功能清单.docx")


def add_frontend_title_block(doc):
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("深海采矿海洋气象预报保障系统")
    set_font(run, size=23, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    sub_run = subtitle.add_run("前端已实现功能清单")
    set_font(sub_run, size=14, bold=True, color="373737")

    add_metadata_rows(doc, [
        ("整理日期", "2026年7月3日"),
        ("项目路径", r"D:\123\0509"),
        ("统计口径", "仅统计前端已实现的页面、组件、交互、地图图层、接口调用封装、前端计算逻辑和静态演示资源。"),
        ("不纳入口径", "不统计 Express 后端服务、数据库表、PostGIS 函数、后端下载脚本和服务端定时任务。"),
    ])
    doc.add_paragraph()


def build_doc():
    doc = Document()
    style_doc(doc)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    hr = header.add_run("前端功能清单 | 深海采矿海洋气象保障")
    set_font(hr, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("仅统计前端实现")
    set_font(fr, size=9, color=MUTED)

    add_frontend_title_block(doc)

    add_note(
        doc,
        "统计结论",
        "当前前端已经形成以 Vue 3 + Cesium 三维地图为核心的业务界面，覆盖矿区总览、采矿系统、环境监测、预警中心、历史数据五类主要工作流，并保留态势总览页签。前端侧已实现的内容包括面板布局、筛选查询、地图交互、图层控制、航线与船舶交互、风险评估展示、历史台风分析界面、ECharts 图表和多类 API 调用封装。"
    )

    add_h(doc, 1, "1. 前端框架与导航")
    add_table(doc, ["功能项", "前端已实现内容", "主要文件"], [
        ("应用框架", "Vue 3 单页应用，集成 Element Plus、Cesium、ECharts、字体资源和统一样式入口。", "src/main.js, package.json"),
        ("顶部页签", "实现“矿区总览、采矿系统、环境监测、预警中心、历史数据、态势总览”六个顶部入口。", "src/constants.js, src/components/Header.vue"),
        ("页签状态切换", "切换不同页签时自动打开默认业务面板，并清理不相关状态，例如航线气象、历史台风轨迹、环境监测图层。", "src/App.vue"),
        ("右侧工具菜单", "按当前页签显示不同工具按钮，支持右侧菜单展开/收起、激活态、高亮态和气象图层子菜单。", "src/components/RightPanel.vue"),
        ("全屏适配", "以 1920 x 1080 为基准对应用容器做缩放适配，窗口尺寸变化时自动更新比例。", "src/App.vue"),
        ("共享视觉样式", "使用科技面板、角标装饰、滚动条、浮动关闭按钮等共享样式类，形成统一前端视觉体系。", "src/styles, src/components/*.vue"),
    ], [1900, 5060, 2400])

    add_h(doc, 1, "2. Cesium 地图与通用地图交互")
    add_table(doc, ["功能项", "前端已实现内容"], [
        ("三维地图容器", "实现 Cesium Viewer 初始化、底图加载、实体管理和地图事件监听。"),
        ("地图工具栏", "实现放大、缩小、复位视角、2D/3D 切换、全屏等常用地图操作。"),
        ("矿区实体加载", "加载矿区 GeoJSON，在地图上展示矿区多边形，并按承包者/矿区属性设置样式。"),
        ("矿区点击交互", "点击矿区后弹出矿区信息，支持展示 ID、承包者、担保国、矿种、位置、合同期限、面积等字段。"),
        ("矿区定位与高亮", "支持从左侧列表、区域查询、历史台风区域和管道模块定位到矿区或区域，并进行地图聚焦/高亮。"),
        ("船舶信息弹窗", "点击船舶或演示船舶后展示船名、类型、船长、船宽、航速、吃水、目的港、ETA、更新时间和航行状态。"),
        ("地图选点", "为点到点航线、矿区科普、避让点、途经点等模块提供地图点选能力。"),
        ("天气点弹窗", "在气象图层或航线气象分析场景中展示点位气象详情。"),
    ], [2500, 6860])

    add_h(doc, 1, "3. 矿区总览前端功能")
    add_table(doc, ["模块", "前端已实现内容"], [
        ("矿区查询分类", [
            "区域查询：区域按钮列表、加载态、选中态。",
            "矿种类型：三类矿种多选。",
            "所属大洋：太平洋、印度洋、大西洋多选。",
            "所属国家：可折叠选择面板、已选标签、单个移除、一键清空。",
            "矿区列表：展示承包者/名称、编号、位置、矿种、担保国和面积，点击联动地图。"
        ]),
        ("矿区图层控制", [
            "太平洋 CCZ 区、太平洋其他区、印度洋区、大西洋区、环境保护区显隐控制。",
            "区域数量显示和定位到该区域。"
        ]),
        ("单矿区总览弹窗", [
            "展示矿区基础信息、水深指标、作业辅助决策、风浪流历史趋势图。",
            "支持“加入气象监测”。"
        ]),
        ("区域总览工作台", [
            "展示区域逐天统计、区域逐 3 小时统计。",
            "展示单矿区逐天统计、单矿区逐 3 小时统计。",
            "用 ECharts 显示风速、浪高、流速、流向、波浪周期等前端图表。"
        ]),
        ("矿区气象监测", [
            "维护监测矿区列表，支持添加、定位、移除。",
            "可配置风速、浪高、洋流阈值。",
            "展示最新气象、预警数量、实时统计卡片、趋势图和预警信息。"
        ]),
        ("矿区科普/作业测算", [
            "支持地图选船位置和目标点。",
            "支持录入船舶载货状态、载重吨、总吨、航速、水深、风速、浪高、流速等参数。",
            "前端计算航线距离、出发/撤退时间、吨位规则匹配和作业结论。"
        ]),
    ], [2200, 7160])

    add_h(doc, 1, "4. 环境监测前端功能")
    add_table(doc, ["功能项", "前端已实现内容", "主要文件"], [
        ("气象图层菜单", "右侧菜单内按分组展示基础气象、OpenWeatherMap、灾害预警等图层，支持组展开、子图层开关、激活数量显示。", "src/components/RightPanel.vue, src/constants.js"),
        ("风场图层", "支持近日风场预报开关，按需加载风场数据并使用 Cesium 风场粒子/图层方式展示。", "src/components/MapContainer.vue, src/utils/wind*"),
        ("海浪图层", "支持近日海浪预报开关，使用风场渲染引擎和热力图展示波浪数据。", "src/components/MapContainer.vue, src/utils/wave*"),
        ("洋流图层", "支持近日洋流预报开关，展示洋流方向/强度并生成热力图。", "src/components/MapContainer.vue, src/utils/oceanCurrent*"),
        ("内波图层", "支持内波图层加载和显隐控制。", "src/components/MapContainer.vue, src/utils/internalWave*"),
        ("OpenWeatherMap 瓦片", "支持云层覆盖、降水分布、温度分布、气压分布四类瓦片叠加。", "src/utils/openWeatherMapLayer.js"),
        ("时间轴", "环境监测页签自动显示时间轴，时间变化后驱动地图更新气象数据。", "src/components/TimelineControl.vue"),
        ("气象点查询", "气象图层开启时可点选地图位置，展示点位气象面板。", "src/components/WeatherPointPicker.vue, src/components/WindyStyleWeatherPanel.vue"),
    ], [2000, 5060, 2300])

    add_h(doc, 1, "5. 采矿系统前端功能")
    add_table(doc, ["模块", "前端已实现内容"], [
        ("管道评估", [
            "矿区选择下拉、矿区加载态和定位联动。",
            "水深、顶端条件、材料、安全系数选择。",
            "当前矿区风速、阵风、浪高、周期、流速、数据时次展示。",
            "生成评估、刷新数据、最终结论、风浪流趋势图、综合应力对比图。"
        ]),
        ("船舶搜索", [
            "单船 MMSI 查询、多船 MMSI 批量查询。",
            "搜索结果卡片、定位按钮、批量结果列表。",
            "已搜索船舶列表和清空功能。"
        ]),
        ("历史轨迹", [
            "MMSI 输入、1 小时/6 小时/24 小时/7 天快捷时间。",
            "自定义开始/结束时间，轨迹查询结果摘要，地图轨迹绘制联动。"
        ]),
        ("航线规划", [
            "港到港和点到点两种规划模式。",
            "点到点支持手输经纬度和地图选点。",
            "高级配置支持避让点、途经点、地图选点、增加和删除节点。",
            "规划结果联动地图绘制航线。"
        ]),
        ("航线气象", [
            "沿航线采样或使用航线点查询气象。",
            "按船速计算各点到达时间。",
            "前端风险阈值可配置，并按安全/注意/警告/危险着色。",
            "气象列表支持刷新、过滤、点击定位。"
        ]),
        ("航线动态演示", [
            "播放、暂停、继续、停止、倍速控制。",
            "演示航线、演示船舶、航点气象、风险警告、矿区气象卡片联动。"
        ]),
        ("区域监控面板", [
            "区域名称、风速阈值、浪高阈值表单。",
            "地图绘制区域、取消绘制、区域列表、显示/隐藏、定位、详情、删除。",
            "区域详情前端对话框展示区域船舶和事件记录。"
        ]),
    ], [2200, 7160])

    add_h(doc, 1, "6. 预警中心前端功能")
    add_table(doc, ["功能项", "前端已实现内容"], [
        ("管道预警入口", "预警中心页签默认打开管道预警面板，并收起右侧菜单。"),
        ("预警列表", "按全部、绿色、黄色、橙色、红色过滤；展示预警消息、矿区、时间、作业类型和处理状态。"),
        ("预警操作", "支持刷新、清空，以及将选中预警标记为已处理或恢复待处理。"),
        ("处置结论", "前端根据预警数据生成处置结论，展示是否可继续作业、是否需要准备或立即撤收。"),
        ("判断依据", "展示预警来源、触发原因、应力判断、海况数据、匹配工况。"),
        ("图表展示", "使用 ECharts 展示风浪流数据图和综合应力图，无数据时显示空状态。"),
        ("全局通知", "前端 WebSocket 消息处理已实现船舶进入、船舶离开、区域预警等通知展示。"),
    ], [2500, 6860])

    add_h(doc, 1, "7. 历史数据前端功能")
    add_table(doc, ["模块", "前端已实现内容"], [
        ("历史台风查询面板", [
            "区域快速切换，显示当前区域和影响范围。",
            "影响范围支持 100 km、300 km、500 km。",
            "影响等级支持全部、核心、强影响、外围。",
            "排序支持距离最近、风速最高、时间最新。"
        ]),
        ("历史台风列表", [
            "展示台风名称、年份/海盆、影响等级、最近距离、最大风速、影响时长、最近时间。",
            "支持分页和选中态。"
        ]),
        ("历史台风窗口工作台", [
            "展示主要窗口、高峰月份、窗口内影响次数。",
            "用 ECharts 展示历史集中窗口月份分布。"
        ]),
        ("当前台风过程", [
            "选中台风后展示进入影响、最近距离、离开影响。",
            "加载轨迹后展示过程图，并联动地图绘制轨迹。"
        ]),
    ], [2200, 7160])

    add_h(doc, 1, "8. 前端计算、状态与服务封装")
    add_table(doc, ["类别", "前端已实现内容", "主要文件"], [
        ("API 封装", "封装矿区、区域总览、台风历史、气象数据、矿区监测等请求；统一基础 URL、超时和端点常量。", "src/api/*.js"),
        ("船讯网前端封装", "封装单船、多船、船队、轨迹、港到港航线、点到点航线、点位气象、海区气象、台风查询等调用。", "src/utils/shipxyApi.js"),
        ("气象数据加载器", "风场、海浪、洋流、内波均支持 API 模式和本地文件模式，按配置动态导入。", "src/config/dataSource.js, src/utils/*Loader*.js"),
        ("风险评估", "前端统一计算蒲福风级、风速/浪高/能见度/涌高/海流综合风险，并输出颜色、描述和风险因素。", "src/utils/weatherRiskAssessment.js"),
        ("管道风险", "前端根据材料、水深、顶端条件、风浪流和阈值计算应力、风险等级、作业建议、作业窗口和预警同步。", "src/utils/pipelineRiskService.js, src/utils/operationRiskDecisionService.js"),
        ("矿区科普测算", "前端计算航行距离、船速、布放/回收耗时、吨位规则和海况等级。", "src/utils/miningScienceService.js"),
        ("本地状态持久化", "管道配置、阈值、预报、风险历史、选型状态和管道预警使用前端存储封装。", "src/utils/pipelineStorage.js, src/utils/pipelineWarningService.js"),
        ("地图图层类", "实现矿区热力图、OpenWeatherMap 图层、船舶图层、航线图层、航线气象图层、航线演示图层、船舶轨迹图层等前端类。", "src/utils/*Layer*.js, src/utils/shipTrajectory.js"),
    ], [1900, 5060, 2400])

    add_h(doc, 1, "9. 前端静态资源与演示能力")
    add_table(doc, ["资源/能力", "前端已实现内容"], [
        ("航线演示数据", "public/route-demo-data.json 用于航线动态演示。"),
        ("点到点航线演示", "public/点到点航线规划.json 当前用于点到点航线规划演示数据。"),
        ("船舶模型", "public/models/ship/cargo_ship 提供货船 glTF 模型、bin 文件和贴图。"),
        ("气象图层图片", "public/image 下包含风、浪、洋流、云、降水、温度、气压、台风等图层相关图片资源。"),
        ("字体与主题", "引入 Noto Sans SC、Rajdhani、Orbitron 字体，并有 base、components、premium-theme、animations、cesium、scrollbar 等样式文件。"),
        ("构建配置", "Vite 集成 Vue 与 Cesium 插件，配置静态资源、二进制文件、开发代理和路径别名。"),
    ], [2600, 6760])

    add_h(doc, 1, "10. 前端口径下的边界说明")
    add_table(doc, ["事项", "说明"], [
        ("不统计后端能力", "本版不列 Express 路由、PostgreSQL/PostGIS 表、Webhook 服务端处理、Copernicus/NOAA 后端服务和数据下载脚本。"),
        ("接口调用不等于服务实现", "src/api 和 src/utils/shipxyApi.js 中的接口封装属于前端已实现；但接口背后的服务端能力不在本版统计范围内。"),
        ("态势总览", "前端已有顶部页签和切换逻辑，但当前未配置专属右侧工具和独立业务工作台。"),
        ("点到点航线", "前端已实现点到点 UI 和地图选点流程；当前数据来源为本地演示 JSON。"),
        ("历史台风与矿区总览", "前端已实现查询、展示、图表和地图联动；数据是否可用取决于外部/独立接口响应。"),
        ("视觉 QA", "当前环境缺少 LibreOffice/soffice 时，无法执行 DOCX 渲染成 PNG 的视觉检查，只能做结构检查。"),
    ], [2700, 6660], header_fill=WARN_FILL)

    add_h(doc, 1, "11. 前端代码依据")
    add_table(doc, ["文件/目录", "说明"], [
        ("src/App.vue", "主状态、页签切换、面板编排、地图联动、WebSocket 消息处理。"),
        ("src/constants.js", "系统标题、页签、右侧工具映射、矿区分区和气象图层配置。"),
        ("src/components", "所有前端业务面板、工作台、弹窗、图表和控件。"),
        ("src/api", "前端 API 调用封装。"),
        ("src/utils", "地图图层、船舶、航线、气象、风险评估、管道评估、矿区科普等前端逻辑。"),
        ("src/styles", "全局样式、主题、组件样式、动画、滚动条和 Cesium 样式。"),
        ("src/config/dataSource.js", "气象数据源 API/本地模式切换。"),
        ("public", "前端静态图片、演示 JSON 和船舶模型资源。"),
        ("vite.config.js, package.json", "前端构建、依赖和开发代理配置。"),
    ], [3000, 6360])

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_doc()
    print(output.resolve())
