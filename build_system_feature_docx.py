from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("系统已实现功能清单.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WARN_FILL = "FFF7E6"
BORDER = "B8C7D9"


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9360")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))


def paragraph_text(cell, text, bold=False, color=None, size=10.5, after=2):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_font(run, size=size, color=color, bold=bold)
    return p


def add_table(doc, headers, rows, widths, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_widths(table, widths)
    set_table_borders(table)

    hdr = table.rows[0].cells
    for idx, title in enumerate(headers):
        set_cell_fill(hdr[idx], header_fill)
        paragraph_text(hdr[idx], title, bold=True, color=INK, size=10.5, after=0)

    for row in rows:
        cells = table.add_row().cells
        for idx, item in enumerate(row):
            if isinstance(item, (list, tuple)):
                first = True
                for line in item:
                    p = cells[idx].paragraphs[0] if first else cells[idx].add_paragraph()
                    first = False
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.25
                    run = p.add_run(str(line))
                    set_font(run, size=10)
            else:
                paragraph_text(cells[idx], str(item), size=10, after=2)
    set_table_widths(table, widths)
    return table


def add_h(doc, level, text):
    return doc.add_heading(text, level=level)


def add_note(doc, title, body, fill=LIGHT_FILL):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    set_table_borders(table, color="D9E2EC")
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_font(r, size=10.5, bold=True, color=INK)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.25
    r2 = p2.add_run(body)
    set_font(r2, size=10)
    return table


def add_metadata_rows(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_widths(table, [1650, 7710])
    set_table_borders(table, color="D5DEE9")
    for row_idx, (label, value) in enumerate(rows):
        label_cell, value_cell = table.rows[row_idx].cells
        set_cell_fill(label_cell, "F2F4F7")
        paragraph_text(label_cell, label, bold=True, color=INK, size=10, after=0)
        paragraph_text(value_cell, value, size=10, after=0)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("深海采矿海洋气象预报保障系统")
    set_font(run, size=23, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("已实现功能清单")
    set_font(r, size=14, color="373737", bold=True)

    add_metadata_rows(doc, [
        ("整理日期", "2026年7月3日"),
        ("项目路径", r"D:\123\0509"),
        ("梳理范围", "当前代码库中的前端组件、后端服务、数据接口、脚本与静态资源。"),
        ("口径说明", "“已实现”表示代码中已有页面、交互、接口封装、服务逻辑或数据处理脚本；外部服务接入和占位能力单独标注。"),
    ])

    doc.add_paragraph()


def build_doc():
    doc = Document()
    style_doc(doc)

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    hr = header.add_run("系统功能清单 | 深海采矿海洋气象保障")
    set_font(hr, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("当前代码库梳理")
    set_font(fr, size=9, color=MUTED)

    add_title_block(doc)

    add_note(
        doc,
        "总体结论",
        "系统已经形成以 Cesium 三维地图为核心的海洋气象保障平台，覆盖矿区总览、环境监测、采矿系统、预警中心和历史数据分析。核心已实现能力包括矿区筛选与统计、风浪流图层、船舶搜索与航线规划、航线气象风险、区域监控、提升管道作业评估、管道预警和历史台风分析。"
    )

    add_h(doc, 1, "1. 系统导航与基础框架")
    add_table(doc, ["功能项", "已实现内容", "代码依据"], [
        ("顶部业务页签", "实现“矿区总览、采矿系统、环境监测、预警中心、历史数据、态势总览”六个入口，并在切换时重置或保留相应面板状态。", "src/constants.js, src/App.vue, src/components/Header.vue"),
        ("右侧工具菜单", "按当前页签显示对应工具按钮，支持展开/收起；气象图层在右侧菜单内提供分组、子图层、激活数量和 TIME 标记。", "src/components/RightPanel.vue"),
        ("全屏响应式容器", "以 1920 x 1080 为基准，对整个应用容器做缩放适配，支持不同屏幕尺寸下的统一布局。", "src/App.vue"),
        ("地图工具栏", "实现放大、缩小、复位视角、2D/3D 切换和全屏操作。", "src/components/MapContainer.vue"),
        ("实时消息", "前端连接 WebSocket，处理船舶进入、离开、区域预警、区域创建和删除等消息，并通过 Element Plus 通知展示。", "src/App.vue, backend/server.js"),
    ], [1900, 5300, 2160])

    add_h(doc, 1, "2. 矿区总览与矿区查询")
    add_table(doc, ["模块", "已实现功能"], [
        ("矿区查询面板", [
            "区域查询：加载区域总览列表，支持区域按钮选择。",
            "矿种筛选：多金属结核、富钴铁锰结壳、多金属硫化物多选。",
            "大洋筛选：太平洋、印度洋、大西洋多选。",
            "国家筛选：可折叠国家选择、已选标签、单个移除和一键清空。",
            "矿区列表：显示承包者、编号、位置、矿种、担保国、面积，点击后联动地图与详情。"
        ]),
        ("矿区图层控制", [
            "按太平洋 CCZ、太平洋其他区、印度洋区、大西洋区、环境保护区等地理分区显隐。",
            "每个分区显示数量标记，并支持定位到区域中心。"
        ]),
        ("单矿区总览弹窗", [
            "展示承包者、担保国、矿种、面积、位置、合同期限和水深指标。",
            "生成作业辅助决策，判断是否继续作业、是否需要关注管道撤收。",
            "展示风速、浪高、流速历史趋势图，支持加入矿区气象监测。"
        ]),
        ("区域总览工作台", [
            "展示区域逐日风浪流统计、区域逐 3 小时统计。",
            "展示单矿区逐日统计和单矿区逐 3 小时统计。",
            "支持点击逐日图切换日期，联动小时级图表。"
        ]),
        ("矿区气象监测", [
            "维护监测矿区列表，支持加入、定位、移除。",
            "支持风速、浪高、洋流阈值设置。",
            "展示最新风浪流数据、预警数量、详细统计卡片、时间序列图和预警信息。"
        ]),
        ("矿区科普/作业测算", [
            "支持地图选船位和目标点，录入船舶载货状态、载重吨、总吨、出发/返航航速。",
            "录入水深、风速、浪高、流速、流向等环境参数。",
            "计算航线距离、出发/撤退时间、吨位规则匹配和可作业结论。"
        ]),
    ], [2100, 7260])

    add_h(doc, 1, "3. 环境监测与气象图层")
    add_table(doc, ["能力", "已实现内容", "备注"], [
        ("基础气象图层", "支持近日风场、近日海浪、近日洋流的开关；地图侧按需加载并使用 Cesium 风场渲染/热力图展示。", "数据源可走 API 或本地二进制文件。"),
        ("OpenWeatherMap 图层", "支持云层覆盖、降水分布、温度分布和气压分布的瓦片图层叠加。", "通过 OpenWeatherMap 瓦片 URL 接入。"),
        ("内波图层", "支持内波数据加载和地图展示，并纳入气象图层开关状态。", "配置数据源为 NASA HRET14。"),
        ("时间轴控制", "当环境监测页签开启时显示时间轴，支持时间索引变化并通知地图刷新风浪流/内波数据。", "src/components/TimelineControl.vue"),
        ("气象点查询", "在气象图层开启时支持地图点选，弹出点位气象详情。", "src/components/WeatherPointPicker.vue"),
        ("图层生命周期", "图层关闭时清理弹窗和实体；切换 2D/3D 时保留或恢复风浪图层显示状态。", "src/components/MapContainer.vue"),
    ], [1900, 5260, 2200])

    add_h(doc, 1, "4. 采矿系统")
    add_table(doc, ["模块", "已实现功能"], [
        ("提升管道安全作业评估", [
            "选择矿区、水深、顶端条件、材料和安全系数。",
            "自动加载矿区风浪流数据，生成当前海况和综合应力评估。",
            "输出最终结论：是否可继续作业、是否需要准备或立即撤收。",
            "使用 ECharts 展示风浪流趋势和综合应力对比，并同步生成管道预警。"
        ]),
        ("船舶搜索", [
            "支持单船 MMSI 查询、多船 MMSI 批量查询、搜索结果定位。",
            "维护已搜索船舶列表，支持点击列表定位和清空。"
        ]),
        ("历史轨迹", [
            "按 MMSI 查询历史轨迹，支持 1 小时、6 小时、24 小时、7 天快捷时间。",
            "支持自定义开始/结束时间，返回轨迹点数并在地图绘制轨迹。"
        ]),
        ("航线规划", [
            "支持港到港模式，输入标准五位港口代码。",
            "支持点到点模式，可手输经纬度或在地图上选点。",
            "支持避让点和途经点高级配置，并在地图绘制航线。"
        ]),
        ("航线气象分析", [
            "沿航线点计算到达时间并查询点位气象。",
            "按风级、浪高、能见度、涌高、海流等阈值计算风险。",
            "在地图上以彩色线段和序号标记展示航线风险，提供气象列表、刷新和过滤。"
        ]),
        ("航线动态演示", [
            "支持演示播放、暂停、继续、停止和倍速控制。",
            "使用演示航线数据、船舶模型、航点气象、风险提醒和到达后矿区气象卡片。"
        ]),
        ("区域监控", [
            "创建监控区域：输入区域名称、风速阈值和浪高阈值，在地图绘制多边形。",
            "支持区域列表、显示/隐藏、定位、详情和删除。",
            "区域详情可查看区域内船舶和事件/预警记录。"
        ]),
    ], [2200, 7160])

    add_h(doc, 1, "5. 预警中心")
    add_table(doc, ["功能项", "已实现内容"], [
        ("管道预警列表", "按全部、绿色、黄色、橙色、红色过滤预警；支持刷新、清空和选择预警记录。"),
        ("预警处置概览", "对选中预警生成处置结论，展示可否继续作业、是否需要管道撤收、风险等级和触发原因。"),
        ("判断依据", "展示预警来源、触发原因、应力判断、海况数据和匹配工况。"),
        ("图表分析", "展示风浪流数据图和综合应力图；无数据时给出空状态。"),
        ("处理状态", "支持将预警标记为已处理或恢复为待处理。"),
        ("实时区域预警", "后端在船舶进入区域、气象超过阈值或定时更新发现风险时创建预警，并通过 WebSocket 推送前端。"),
    ], [2200, 7160])

    add_h(doc, 1, "6. 历史数据")
    add_table(doc, ["模块", "已实现功能"], [
        ("历史台风查询", [
            "按区域快速切换历史台风查询范围。",
            "支持 100 km、300 km、500 km 影响范围切换。",
            "支持全部、核心、强影响、外围等级筛选。",
            "支持按最近距离、最大风速、最新时间排序。"
        ]),
        ("历史台风列表", [
            "展示台风名称、年份/海盆、影响等级、最近距离、最大风速、影响时长和最近时间。",
            "支持分页。"
        ]),
        ("历史台风窗口", [
            "统计历史集中月份、峰值月份、窗口内影响次数。",
            "用图表展示历史集中窗口与其他月份分布。"
        ]),
        ("台风轨迹过程", [
            "选中台风后加载轨迹，展示进入影响、最近距离、离开影响。",
            "在地图上绘制台风轨迹，并突出影响范围内的关键段。"
        ]),
    ], [2200, 7160])

    add_h(doc, 1, "7. 后端服务与数据库")
    add_table(doc, ["类别", "已实现内容", "主要文件"], [
        ("区域监控后端", "Express 服务实现监控区域创建、列表、删除、区域船舶、事件日志、Webhook 接收、测试事件和 WebSocket 推送。", "backend/server.js"),
        ("船舶事件处理", "处理船舶进入/离开区域，查询船舶详情和点位气象，进行风速/浪高风险评估，写入事件日志和预警。", "backend/server.js"),
        ("定时更新", "每 10 分钟更新活跃区域内船舶状态、最新气象和风险等级，并广播更新。", "backend/server.js"),
        ("PostgreSQL/PostGIS", "定义监控区域、区域内船舶、预警记录、事件日志、系统配置表；建立空间索引、JSONB 索引、统计视图和空间函数。", "backend/database_postgres.sql"),
        ("Copernicus 数据服务", "提供状态查询、手动更新和区域气象查询；区域查询按多边形中心点读取波高和洋流二进制数据并估算风速。", "backend/routes/copernicusRoutes.js, backend/services/copernicusDataService.js"),
        ("矿区总览占位接口", "提供 /api/mining-areas/:id/overview，占位返回水深和历史风浪流结构，便于前端联调。", "backend/routes/miningAreaOverviewRoutes.js"),
        ("NOAA 服务模块", "代码中已有 NOAA/GFS 风、浪、流下载、GRIB2 转换、缓存和定时更新服务及路由模块。", "backend/services/noaaDataService.js, backend/routes/noaaRoutes.js"),
    ], [1900, 5160, 2300])

    add_h(doc, 1, "8. 数据源、外部服务与脚本")
    add_table(doc, ["数据/服务", "已实现接入或支撑能力"], [
        ("船讯网 ShipXY", "前端封装单船、多船、历史轨迹、港到港航线、点位气象、海区气象、台风列表/详情、区域监控 API；开发环境通过 Vite 代理。"),
        ("矿区风浪流总览服务", "前端已接入 /api/mining-overview/regions、daily、hourly、sites、site daily/hourly，用于区域和站点风浪流统计。"),
        ("气象数据服务", "前端封装 /api/weather/metadata、data、binary、available、point-query；风、浪、流、内波加载器支持 API 与本地文件两种模式。"),
        ("OpenWeatherMap", "接入云、降水、温度、气压瓦片图层。"),
        ("Copernicus/NOAA 数据脚本", "包含 Copernicus 下载转换、GFS 风场下载、CDS 风场下载、GRIB2 转 JSON、二进制转换、洋流降采样等脚本。"),
        ("静态演示资源", "包含点到点航线规划 JSON、航线演示数据、天气图层图片、货船 3D 模型与贴图。"),
        ("构建与前端技术栈", "Vue 3、Vite、Cesium、Element Plus、ECharts、Tailwind/PostCSS、vite-plugin-cesium。"),
    ], [2500, 6860])

    add_h(doc, 1, "9. 主要接口清单")
    add_table(doc, ["接口或入口", "能力", "归属/状态"], [
        ("/api/areas", "创建和获取监控区域。", "本仓库后端实现"),
        ("/api/areas/:id", "删除监控区域。", "本仓库后端实现"),
        ("/api/areas/:id/ships", "获取区域内船舶及船舶详情、最近气象、风险等级。", "本仓库后端实现"),
        ("/api/areas/:id/events", "获取区域事件日志和预警记录。", "本仓库后端实现"),
        ("/webhook/area", "接收船讯网区域进入/离开事件推送。", "本仓库后端实现"),
        ("/api/copernicus/status/update/query-area", "Copernicus 数据状态、后台更新、区域风浪流查询。", "本仓库后端实现"),
        ("/api/mining-areas/:id/overview", "矿区总览占位数据。", "本仓库后端实现，当前为占位结构"),
        ("/api/mining-overview/*", "区域/站点风浪流和海深统计。", "前端已接入，Vite 代理到独立服务"),
        ("/api/weather/*", "风、浪、流、内波数据和点查询。", "前端已接入，Vite 代理到后端服务"),
        ("/api/mining-regions/:id/typhoon/events", "区域历史台风事件查询。", "前端已接入，外部/独立服务"),
        ("/api/typhoons/:sid/track", "台风轨迹查询。", "前端已接入，外部/独立服务"),
        ("/api/shipxy/apicall/v3/*", "船舶、轨迹、航线、气象、台风等船讯网 API。", "开发环境 Vite 代理，生产环境直连"),
    ], [3000, 4200, 2160])

    add_h(doc, 1, "10. 当前边界与需继续完善项")
    add_table(doc, ["事项", "当前状态"], [
        ("态势总览", "顶部页签已实现，但当前 constants 中未配置专属右侧工具，进入后主要表现为清空/关闭业务面板的占位页签。"),
        ("点到点航线规划", "当前使用 public/点到点航线规划.json 本地演示数据；港到港规划走船讯网接口。"),
        ("矿区单点总览后端", "本仓库 /api/mining-areas/:id/overview 已实现占位结构；真实水深和历史风浪流主要依赖 /api/mining-overview 独立服务。"),
        ("NOAA 路由", "NOAA 服务和路由模块代码已存在，但当前 backend/server.js 未注册 /api/noaa 路由。"),
        ("灾害预警图层", "气象图层配置包含台风路径预警、海啸传播预警、内波；当前地图逻辑显式实现了内波加载，历史台风已作为独立模块实现。"),
        ("外部服务依赖", "矿区总览、历史台风、天气数据等多个接口通过 Vite 代理或生产地址指向独立服务，部署时需要保证这些服务可用。"),
    ], [2900, 6460], header_fill=WARN_FILL)

    add_h(doc, 1, "11. 梳理依据")
    add_table(doc, ["文件/目录", "用途"], [
        ("src/App.vue, src/constants.js", "主流程、页签、面板状态、功能映射和全局事件。"),
        ("src/components", "各业务面板、工作台、弹窗、图表和交互控件。"),
        ("src/api, src/utils", "接口封装、数据加载器、地图图层、风险评估、管道评估、航线与船舶服务。"),
        ("backend/server.js, backend/routes, backend/services", "Express API、Webhook、WebSocket、Copernicus/NOAA 服务模块。"),
        ("backend/database_postgres.sql", "区域监控数据库表、索引、视图和空间函数。"),
        ("vite.config.js, package.json", "前端构建、开发代理和依赖栈。"),
        ("public, backend/scripts", "静态演示数据、模型、图片和数据下载/转换脚本。"),
    ], [3000, 6360])

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_doc()
    print(output.resolve())
